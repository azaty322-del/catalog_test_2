"""
Импорт каталога UTV Media из Excel в data/data.json.

Источник заполняется людьми, поэтому одно и то же поле приходит в разных форматах:
длительность бывает числом часов, time и timedelta, год и число эпизодов —
числом или диапазоном в строке. Все разборы собраны здесь, чтобы приложение
работало с уже нормализованными данными.

Запуск: python3 scripts/import-catalog.py путь-к-файлу.xlsx
"""

import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

# Порядок листов задаёт порядок разделов в навигации сайта.
CATEGORIES = [
    ("Природа", "nature", "Nature", True),
    ("Путешествия", "travel", "Travel", True),
    ("Кулинария и еда", "cooking", "Food & Cooking", True),
    ("Научпоп", "science", "Popular Science", False),
    ("Познавательный", "educational", "Educational", False),
    ("Эксперты", "experts", "Experts", False),
    ("Дача и сад", "garden", "Garden & Country", False),
    ("Детям и подросткам", "kids", "Kids & Teens", False),
    ("Тренировки", "fitness", "Fitness", False),
    ("Под запрос", "on-request", "On Request", False),
]

TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
    "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
    "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "ts",
    "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
    "я": "ya",
}


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def slugify(text: str) -> str:
    out = "".join(TRANSLIT.get(ch, ch) for ch in text.lower())
    out = re.sub(r"[^a-z0-9]+", "-", out)
    return re.sub(r"-{2,}", "-", out).strip("-")[:60] or "item"


def parse_int_range(value):
    """Возвращает (min, max). Принимает 30, «15-23», «18—22»."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        n = int(value)
        return (n, n)
    numbers = [int(n) for n in re.findall(r"\d+", str(value))]
    if not numbers:
        return None
    return (min(numbers), max(numbers))


def parse_years(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        year = int(value)
        return (year, year)
    years = [int(y) for y in re.findall(r"\b(19|20)\d{2}\b", str(value))] or [
        int(y) for y in re.findall(r"\d{4}", str(value))
    ]
    years = [int(y) for y in re.findall(r"\d{4}", str(value))]
    if not years:
        return None
    return (min(years), max(years))


def parse_total_minutes(value):
    """Общая длительность: число часов, time, timedelta или строка «ч:мм:сс»."""
    if value is None:
        return None
    if isinstance(value, datetime.timedelta):
        return round(value.total_seconds() / 60)
    if isinstance(value, datetime.time):
        return value.hour * 60 + value.minute + round(value.second / 60)
    if isinstance(value, datetime.datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, (int, float)):
        return round(float(value) * 60)

    text = clean(value)
    if ":" in text:
        parts = [int(p) for p in re.findall(r"\d+", text)]
        hours = parts[0] if parts else 0
        minutes = parts[1] if len(parts) > 1 else 0
        return hours * 60 + minutes
    numbers = re.findall(r"[\d.,]+", text)
    return round(float(numbers[0].replace(",", ".")) * 60) if numbers else None


def parse_episode_duration(value):
    if isinstance(value, datetime.datetime):
        return (value.hour * 60 + value.minute, value.hour * 60 + value.minute)
    return parse_int_range(value)


def split_multi(value: str):
    """«Сербия / Россия», «HD / 4K» — несколько значений в одной ячейке."""
    return [part.strip() for part in re.split(r"[/,]", value) if part.strip()]


# «Съёмки со всего мира» описывает способ производства, а не страну.
COUNTRY_ALIASES = {
    "съемки со всего мира": "Весь мир",
    "съёмки со всего мира": "Весь мир",
}

LANGUAGE_CODES = {"русский": "ru", "английский": "en", "russian": "ru", "english": "en"}


def detect_host(url: str):
    if not url:
        return None
    if "youtu" in url:
        return "youtube"
    if "rutube" in url:
        return "rutube"
    if "disk.yandex" in url or "yadi.sk" in url:
        return "yandex-disk"
    if "vk.com" in url or "vkvideo" in url:
        return "vk"
    return "other"


def localized(ru: str):
    return {"ru": ru, "en": None}


def main(source: Path, output: Path) -> None:
    wb = openpyxl.load_workbook(source, data_only=True)

    categories, movies, skipped = [], [], []
    used_slugs = set()

    for index, (sheet_name, slug, title_en, in_nav) in enumerate(CATEGORIES):
        if sheet_name not in wb.sheetnames:
            print(f"  лист «{sheet_name}» не найден — пропущен")
            continue

        rows = [
            row
            for row in wb[sheet_name].iter_rows(min_row=2, values_only=True)
            if row and row[0] and clean(row[0])
        ]

        categories.append(
            {
                "id": f"cat-{index + 1}",
                "slug": slug,
                "title": {"ru": sheet_name, "en": title_en},
                "showInNav": in_nav,
                "isEmpty": len(rows) == 0,
            }
        )

        for row in rows:
            title = clean(row[0])

            years = parse_years(row[7])
            duration = parse_episode_duration(row[3])
            episodes = parse_int_range(row[2])
            total = parse_total_minutes(row[4])

            # Без года или длительности карточку показать нечем — откладываем.
            if years is None or duration is None:
                skipped.append((sheet_name, title, "нет года или длительности"))
                continue

            base_slug = slugify(title)
            movie_slug = base_slug
            suffix = 2
            while movie_slug in used_slugs:
                movie_slug = f"{base_slug}-{suffix}"
                suffix += 1
            used_slugs.add(movie_slug)

            country_raw = clean(row[6])
            countries = [
                COUNTRY_ALIASES.get(part.lower(), part) for part in split_multi(country_raw)
            ]

            languages = [
                LANGUAGE_CODES.get(part.lower(), part)
                for part in split_multi(clean(row[8]))
            ]

            url = clean(row[11])
            tags_raw = clean(row[14])
            rights = clean(row[13])

            movies.append(
                {
                    "id": f"mv-{len(movies) + 1:03d}",
                    "sourceId": clean(row[15]) or f"{slug}-{len(movies) + 1}",
                    "slug": movie_slug,
                    "category": slug,
                    "title": localized(title),
                    "description": localized(clean(row[1])),
                    "episodes": episodes[0] if episodes else 1,
                    "episodeDuration": {"min": duration[0], "max": duration[1]},
                    "totalDurationMinutes": total
                    or (episodes[0] if episodes else 1) * duration[0],
                    "qualities": split_multi(clean(row[5])),
                    "countries": countries,
                    "yearFrom": years[0],
                    "yearTo": years[1],
                    "languages": languages,
                    "author": clean(row[9]),
                    "ageRating": clean(row[10]),
                    "distributionTerritory": clean(row[12]),
                    "rightsPeriod": rights if rights and rights != "-" else None,
                    "tags": []
                    if not tags_raw or tags_raw == "-"
                    else [t.strip() for t in tags_raw.split(",") if t.strip()],
                    "trailer": {"url": url or None, "host": detect_host(url)},
                    "poster": "/placeholder.png",
                    "stills": [],
                }
            )

    payload = {
        "meta": {
            "source": source.name,
            "note": "Сгенерировано scripts/import-catalog.py. Правки вносите в Excel и запускайте импорт заново.",
        },
        "categories": categories,
        "movies": movies,
    }

    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"разделов: {len(categories)}, позиций: {len(movies)}")
    for category in categories:
        count = sum(1 for m in movies if m["category"] == category["slug"])
        print(f"  {category['title']['ru']:24} {count:>3}")
    if skipped:
        print("\nне импортировано:")
        for sheet, title, reason in skipped:
            print(f"  {sheet}: {title[:45]} — {reason}")


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Каталог_контента_UTV_Media.xlsx")
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("data/data.json")
    main(src, dst)
